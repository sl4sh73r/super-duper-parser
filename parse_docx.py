import re  # Регулярные выражения для поиска шаблонов в тексте
import docx  # Работа с DOCX файлами
import fitz  # PyMuPDF для работы с PDF файлами
from bs4 import BeautifulSoup  # Для парсинга HTML
from datetime import datetime, timedelta  # Работа с датами и временем
import pandas as pd  # Работа с данными

# Функция для извлечения данных из таблицы в DOCX файле
def extract_data_from_table(docx_path):
    doc = docx.Document(docx_path)  # Открываем DOCX файл
    for table in doc.tables:  # Проходим по всем таблицам в документе
        first_cell_text = table.cell(0, 0).text  # Извлекаем текст из первой ячейки
        if "УСЛУГИ ПО ПРЕДОСТАВЛЕНИЮ ДОСТУПА К ПРОГРАММНОМУ ПРОДУКТУ" in first_cell_text:
            right_cell_text = table.cell(3, 3).text  # Извлекаем текст из ячейки (3, 3)
            dates = find_dates_in_section(right_cell_text)  # Ищем даты в тексте ячейки
            if dates:
                return dates  # Возвращаем найденные даты
    return None  # Если ничего не найдено, возвращаем None

# Функция для поиска дат в секции текста
def find_dates_in_section(section):
    section = section.replace('\n', ' ')  # Заменяем переносы строк на пробелы
    date_pattern = re.compile(r'\b\d{2}\.\d{2}\.\d{4}\b')  # Шаблон для поиска дат в формате DD.MM.YYYY
    dates = date_pattern.findall(section)  # Ищем все даты в секции
    return dates  # Возвращаем найденные даты

# Функция для извлечения текста из PDF файла
def extract_text_from_pdf(pdf_path):
    document = fitz.open(pdf_path)  # Открываем PDF файл
    text = ""
    for page_num in range(len(document)):  # Проходим по всем страницам документа
        page = document.load_page(page_num)  # Загружаем страницу
        text += page.get_text()  # Извлекаем текст со страницы
    return text  # Возвращаем весь текст

# Функция для поиска секции текста между двумя шаблонами
def find_section(text, start_pattern, end_pattern):
    pattern = re.compile(rf'({start_pattern}.*?){end_pattern}', re.DOTALL)  # Компилируем регулярное выражение
    match = pattern.search(text)  # Ищем совпадение в тексте
    if match:
        return match.group(0)  # Возвращаем найденную секцию
    return None  # Если ничего не найдено, возвращаем None

# Функция для поиска периода оплаты в тексте
def find_payment_period(text):
    sections_to_check = ['2\\.6\\.2', '2\\.7\\.2']  # Секции для проверки
    keywords = ['по факту', 'ежемесячно', 'ежеквартально']  # Ключевые слова для поиска
    found_keywords = set()  # Множество для хранения найденных ключевых слов
    for section in sections_to_check:  # Проходим по секциям
        section_text = find_section(text, section, '\\n\\s*\\n')  # Ищем текст секции
        if section_text:
            for keyword in keywords:  # Проходим по ключевым словам
                if keyword in section_text:
                    found_keywords.add(keyword)  # Добавляем найденное ключевое слово в множество
    return found_keywords  # Возвращаем найденные ключевые слова

# Функция для поиска номера статьи 4 в тексте
def find_article_4_number(text):
    section = find_section(text, '4\\.1', '\\n\\s*\\н')  # Ищем секцию текста
    if section:
        section = section.replace('4.1', '', 1)  # Убираем '4.1' один раз из начала
        match = re.search(r'\b(\d+)\b', section)  # Ищем число в секции
        if match:
            return int(match.group(1))  # Преобразуем в число и возвращаем
    return None  # Если ничего не найдено, возвращаем None

# Функция для извлечения даты из HTML файла
def extract_date_from_html(html_path):
    with open(html_path, 'r', encoding='utf-8') as file:  # Открываем HTML файл
        soup = BeautifulSoup(file, 'html.parser')  # Парсим HTML
        tables = soup.find_all('table')  # Ищем все таблицы
        for table in tables:  # Проходим по всем таблицам
            cell = table.find('td', string="Дата и время подписания ")  # Ищем ячейку с нужным текстом
            if cell:
                rows = table.find_all('tr')  # Ищем все строки таблицы
                if len(rows) >= 3:
                    columns = rows[2].find_all('td')  # Ищем все ячейки в третьей строке
                    if len(columns) >= 4:
                        date_cell = columns[3]  # Извлекаем ячейку с датой
                        return date_cell.text[:10]  # Оставляем только первые 10 символов
    return None  # Если ничего не найдено, возвращаем None

# Функция для парсинга нескольких HTML файлов
def parse_multiple_html(html_paths):
    html_dates = []  # Список для хранения дат
    for path in html_paths:  # Проходим по всем путям к HTML файлам
        date = extract_date_from_html(path)  # Извлекаем дату из HTML файла
        if date:
            html_dates.append(date)  # Добавляем дату в список
    return html_dates  # Возвращаем все найденные даты

# Функция для получения первого рабочего дня месяца
def get_first_working_day(year, month):
    first_day = datetime(year, month, 1)  # Получаем первый день месяца
    while first_day.weekday() > 4:  # Проверяем, является ли первый день месяца рабочим днем (0 - понедельник, 4 - пятница)
        first_day += timedelta(days=1)  # Если нет, переходим к следующему дню
    return first_day  # Возвращаем первый рабочий день

# Функция для добавления рабочих дней к дате
def add_business_days(start_date, business_days):
    current_date = start_date  # Начальная дата
    days_added = 0  # Количество добавленных дней
    while days_added < business_days-1:
        current_date += timedelta(days=1)  # Добавляем один день
        if current_date.weekday() < 5:  # Если это рабочий день (0 - понедельник, 4 - пятница)
            days_added += 1  # Увеличиваем количество добавленных дней
    return current_date  # Возвращаем конечную дату

# Функция для расчета крайних сроков на основе номера статьи 4 и дат из HTML файлов
def calculate_deadline_dates(article_4_number, html_dates):
    deadlines = []  # Список для хранения крайних сроков
    for date_str in html_dates:  # Проходим по всем датам из HTML файлов
        year = int(date_str[-4:])  # Извлекаем год из строки даты
        month = int(date_str[3:5])  # Извлекаем месяц из строки даты
        first_working_day = get_first_working_day(year, month)  # Получаем первый рабочий день месяца
        deadline_date = add_business_days(first_working_day, article_4_number)  # Добавляем рабочие дни к дате
        deadlines.append(deadline_date.strftime('%d.%m.%Y'))  # Форматируем дату и добавляем в список
    return deadlines  # Возвращаем все крайние сроки

# Функция для парсинга DOCX и PDF файлов
def parse_docx_and_pdf(docx_path, pdf_path):
    docx_dates = extract_data_from_table(docx_path)  # Извлекаем данные из таблицы в DOCX файле
    if docx_dates:
        pdf_text = extract_text_from_pdf(pdf_path)  # Извлекаем текст из PDF файла
        payment_period = find_payment_period(pdf_text)  # Ищем период оплаты в тексте PDF файла
        article_4_number = find_article_4_number(pdf_text)  # Ищем номер статьи 4 в тексте PDF файла
        return docx_dates, payment_period, article_4_number  # Возвращаем найденные данные
    return None, None, None  # Если ничего не найдено, возвращаем None